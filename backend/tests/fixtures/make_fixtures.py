"""Builds the client-style timetables used by tests/test_client_requirements.py.

Topics are deliberately scattered across days and weeks, the way coordinators upload them:
2nd Year MBBS (.docx grid), 4th Year MBBS (.xlsx, one sheet per week), 1st Year BDS (pipe text).
Run from backend/: python tests/fixtures/make_fixtures.py
"""
import docx
from openpyxl import Workbook
from pathlib import Path

OUT = Path(__file__).parent
HDR = ["DAYS", "08:00 – 09:00", "09:00 – 10:00", "10:00 – 11:00", "11:00 – 12:00", "12:00 – 12:30", "12:30 – 02:00"]
DAYS = ["MONDAY", "TUESDAY", "WEDNESDAY", "THURSDAY", "FRIDAY", "SATURDAY"]

def L(disc, title, fac, kind="LGIS"):
    return f"{kind}\n{disc}\n{title}\n{fac}"

# ---- 2nd Year MBBS: anatomy/physiology/biochem of the same systems scattered over 2 weeks ----
mbbs2_weeks = [
    {  # week 1 (random order on purpose)
        "MONDAY":    [L("ANATOMY", "Heart chambers and valves", "Dr. Ambar Saba"), L("BIOCHEM", "Renal function tests", "Dr. Asma Khan"), L("PHYSIO", "Nerve conduction and action potential", "Dr. Jillani"), L("HISTO", "Histology of lymph node", "Dr. Gul Rukh")],
        "TUESDAY":   [L("ANATOMY", "Brachial plexus", "Dr. Haris Khan"), L("PHYSIO", "Mechanics of breathing", "Dr. Javied Iqbal"), L("BIOCHEM", "Glucose metabolism and glycolysis", "Dr. Asma Khan"), L("EMBRYO", "Development of kidney", "Dr. Ambar Saba")],
        "WEDNESDAY": [L("PHYSIO", "Renal filtration and GFR", "Dr. Jillani"), L("ANATOMY", "Coronary arteries", "Dr. Ambar Saba"), L("BIOCHEM", "Vitamin D and calcium metabolism", "Dr. Arsalan"), L("HISTO", "Histology of cardiac muscle", "Dr. Gul Rukh")],
        "THURSDAY":  [L("ANATOMY", "Bronchopulmonary segments of lung", "Dr. Haris Khan"), L("PHYSIO", "Cardiac cycle", "Dr. Javied Iqbal"), "", ""],
        "FRIDAY":    [L("BIOCHEM", "Lipoprotein metabolism and cholesterol", "Dr. Arsalan"), L("HISTO", "Histology of kidney and nephron", "Dr. Gul Rukh"), L("PHYSIO", "Regulation of blood pressure", "Dr. Jillani"), ""],
        "SATURDAY":  [L("ANATOMY", "Kidney gross structure and relations", "Dr. Haris Khan"), L("EMBRYO", "Development of heart", "Dr. Ambar Saba"), "", ""],
    },
    {  # week 2
        "MONDAY":    [L("PHYSIO", "Cardiac cycle", "Dr. Javied Iqbal"), L("ANATOMY", "Axilla and its contents", "Dr. Haris Khan"), L("BIOCHEM", "Cardiac enzymes in myocardial infarction", "Dr. Asma Khan"), L("PHYSIO", "Gas exchange in lungs", "Dr. Javied Iqbal")],
        "TUESDAY":   [L("HISTO", "Histology of peripheral nerve", "Dr. Gul Rukh"), L("PHYSIO", "ECG and conducting system of heart", "Dr. Jillani"), "", ""],
        "WEDNESDAY": [L("ANATOMY", "Pleura and its recesses", "Dr. Haris Khan"), L("PHYSIO", "Tubular reabsorption and urine concentration", "Dr. Jillani"), "", ""],
        "THURSDAY":  [L("BIOCHEM", "Haemoglobin structure and function", "Dr. Arsalan"), L("PHYSIO", "Erythropoiesis and anaemia", "Dr. Javied Iqbal"), "", ""],
        "FRIDAY":    [L("ANATOMY", "Ureter and urinary bladder", "Dr. Ambar Saba"), "", "", ""],
        "SATURDAY":  [L("MENTOR MENTEE MEETING", "Mentor mentee meeting", ""), "", "", ""],
    },
]

def build_docx(path, year_line, module_line_fmt, weeks):
    d = docx.Document()
    for w_idx, week in enumerate(weeks, 1):
        d.add_paragraph(year_line)
        d.add_paragraph(module_line_fmt.format(w=w_idx))
        t = d.add_table(rows=1 + len(DAYS), cols=len(HDR))
        for c, h in enumerate(HDR):
            t.rows[0].cells[c].text = h
        for r, day in enumerate(DAYS, 1):
            row = t.rows[r].cells
            row[0].text = day
            slots = week.get(day, ["", "", "", ""])
            for c, txt in enumerate(slots, 1):
                row[c].text = txt
            row[5].text = "HALF DAY" if day == "FRIDAY" else "BREAK"
            row[6].text = "" if day == "FRIDAY" else "SGD\nBatch A: Anat\nBatch B: Phy\nBatch C: Bio\nBatch D: SDL"
        d.add_paragraph("HISTO P1= Cardiac muscle   PHYSIO P1= Blood pressure measurement")
        d.add_paragraph("BIOCHEM P1= Estimation of blood glucose   SKILL LAB= Cardiovascular examination")
        if w_idx < len(weeks):
            d.add_page_break()
    d.save(path)

build_docx(OUT / "MBBS_Year2_CardioRenal.docx", "TIMETABLE FOR 2ND YEAR MBBS",
           "CARDIO-RENAL BLOCK-II (MODULE-04) (WEEK-0{w}) (SESSION 2026-2027)", mbbs2_weeks)

# ---- 4th Year MBBS (xlsx, one sheet per week): eye / medicine / surgery on the same diseases ----
mbbs4_weeks = [
    {
        "MONDAY":    [L("CLINICAL FACULTY EYE", "Diabetic retinopathy", "Dr. Nasir Mahmood"), L("SURGERY", "Acute appendicitis", "Dr. Ayyaz Afridi"), L("MEDICINE", "Hypertension: diagnosis and management", "Dr. M. Arshad"), L("ENT", "Otitis media", "Dr. Saima Gul")],
        "TUESDAY":   [L("SURGERY", "Thyroid swelling and goitre: surgical management", "Dr. Ayyaz Afridi"), L("CLINICAL FACULTY EYE", "Primary open angle glaucoma", "Dr. Nasir Mahmood"), L("PEADS", "Type 1 diabetes in children", "Dr. Farah Naz"), L("MEDICINE", "Chronic kidney disease", "Dr. M. Arshad")],
        "WEDNESDAY": [L("CLINICAL FACULTY EYE", "Cataract", "Dr. Nasir Mahmood"), L("MEDICINE", "Diabetes mellitus: diagnosis and management", "Dr. M. Arshad"), L("ENT", "Epistaxis", "Dr. Saima Gul"), ""],
        "THURSDAY":  [L("SURGERY", "Diabetic foot", "Dr. Ayyaz Afridi"), L("MEDICINE", "Thyroid disorders: hyper and hypothyroidism", "Dr. M. Arshad"), "", ""],
        "FRIDAY":    [L("CLINICAL FACULTY EYE", "Hypertensive retinopathy", "Dr. Nasir Mahmood"), L("SURGERY", "Peripheral arterial disease and atherosclerosis", "Dr. Ayyaz Afridi"), "", ""],
        "SATURDAY":  [L("MEDICINE", "Ischaemic heart disease and angina", "Dr. M. Arshad"), L("PEADS", "Acute glomerulonephritis in children", "Dr. Farah Naz"), "", ""],
    },
]
wb = Workbook(); wb.remove(wb.active)
for w_idx, week in enumerate(mbbs4_weeks, 1):
    ws = wb.create_sheet(f"Week {w_idx}")
    ws.append(["TIMETABLE FOR 4TH YEAR MBBS"])
    ws.append([f"SPECIAL SENSES & ENDOCRINE BLOCK (MODULE-02) (WEEK-0{w_idx}) (SESSION 2026-2027)"])
    ws.append(HDR)
    for day in DAYS:
        slots = week.get(day, ["", "", "", ""])
        ws.append([day, *slots, "HALF DAY" if day == "FRIDAY" else "BREAK", ""])
wb.save(OUT / "MBBS_Year4_SpecialSenses.xlsx")

# ---- 1st Year BDS (pasted/pipe text) ----
(OUT / "BDS_Year1_HeadNeck.txt").write_text("""TIMETABLE FOR 1ST YEAR BDS
HEAD & NECK BLOCK-I (MODULE-01) (SESSION 2026-2027)

Oral Biology | LGIS | Tooth development and odontogenesis | Dr. Sana Ullah
Physiology | LGIS | Saliva composition and secretion | Dr. Jillani
Anatomy | LGIS | Temporomandibular joint | Dr. Haris Khan
Biochemistry | LGIS | Calcium and phosphate metabolism | Dr. Asma Khan
Anatomy | LGIS | Muscles of mastication | Dr. Haris Khan
Oral Biology | LGIS | Enamel structure and amelogenesis | Dr. Sana Ullah
Physiology | LGIS | Mastication and swallowing | Dr. Jillani
Anatomy | LGIS | Parotid and submandibular salivary glands | Dr. Ambar Saba
Dental Materials | LGIS | Gypsum products | Dr. Bilal Ahmed
Oral Biology | LGIS | Dentin and pulp complex | Dr. Sana Ullah
Anatomy | LGIS | Facial nerve | Dr. Haris Khan
Physiology | LGIS | Blood groups and transfusion | Dr. Javied Iqbal
""", encoding="utf-8")
print("written:", sorted(p.name for p in OUT.iterdir() if p.suffix in {".docx", ".xlsx", ".txt"}))
