"""Build the anonymised real-timetable fixture used by tests/test_real_timetable.py.

Source: docs/sample_timetable.docx (the client's 2nd Year MBBS Neurosciences module, 7 weeks).
Every teacher name (all spelling variants of one person) becomes one code, 'Dr. T01', 'Dr. T02', ...
in order of first appearance; titles are normalised to 'Dr.'. Nothing else in the file changes, so
the fixture has the exact structure of the real timetable (merged cells, time rows, legends).

    python tests/fixtures/make_real_fixture.py
"""
from __future__ import annotations

import io
import re
import sys
from pathlib import Path

import docx

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path.insert(0, str(HERE.parents[1]))

from app.services import parser  # noqa: E402
from app.services.timetable_grid import _NAME_STOP, _TITLE_CANON, _TITLE_RE  # noqa: E402

SOURCE = ROOT / "docs" / "sample_timetable.docx"
TARGET = HERE / "real_neuro_module07.docx"


def name_spans(text: str) -> list[tuple[int, int, str]]:
    """(start, end, normalised name) of every 'Dr. Name' in a paragraph (same rules as the parser)."""
    spans = []
    for m in _TITLE_RE.finditer(text):
        if m.start() > 0 and text[m.start() - 1].isalpha():
            continue
        rest_start = m.end()
        rest = re.split(r"[/()\n,;:\\|]|\s[-–]\s", text[rest_start:], maxsplit=1)[0]
        tokens, end = [], rest_start
        for tm in re.finditer(r"[A-Za-z][A-Za-z0-9]*\.?", rest):
            word = tm.group(0).rstrip(".")
            if word.lower() in _NAME_STOP or len(tokens) >= 3:
                break
            if not tokens and not word[0].isupper():
                break
            if len(word) >= 3 and word.isupper() and tokens:
                break
            if word.lower() in _TITLE_CANON:
                break
            tokens.append(word)
            end = rest_start + tm.end()
        if tokens:
            parts = [f"{t.upper()}." if len(t) == 1 else t[0].upper() + t[1:].lower() for t in tokens]
            spans.append((m.start(), end, f"{_TITLE_CANON[m.group(1).lower()]} {' '.join(parts)}"))
    return spans


def main() -> None:
    data = SOURCE.read_bytes()
    secs = parser.read_upload_sections("x.docx", data)
    res = parser.parse_sections([(d, "x.docx") for d in secs], use_llm=False)
    aliases = res.teacher_aliases
    codes: dict[str, str] = {}

    def code_for(name: str) -> str:
        canon = aliases.get(name, name)
        if canon not in codes:
            codes[canon] = f"Dr. T{len(codes) + 1:02d}"
        return codes[canon]

    doc = docx.Document(io.BytesIO(data))
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        seen = set()
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue
                seen.add(id(cell._tc))
                paragraphs.extend(cell.paragraphs)
    changed = 0
    for p in paragraphs:
        text = p.text
        spans = name_spans(text)
        if not spans:
            continue
        out, last = [], 0
        for s, e, name in spans:
            out.append(text[last:s])
            out.append(code_for(name))
            last = e
        out.append(text[last:])
        new = "".join(out)
        if p.runs:
            p.runs[0].text = new
            for r in p.runs[1:]:
                r.text = ""
        changed += 1
    doc.save(TARGET)
    print(f"wrote {TARGET.name}: {changed} paragraphs, {len(codes)} teachers coded")


if __name__ == "__main__":
    main()
